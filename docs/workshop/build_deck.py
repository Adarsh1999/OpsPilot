"""Build the workshop from the supplied Microsoft template, keeping template parts intact."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from zipfile import ZipFile, ZIP_DEFLATED
from io import BytesIO
import hashlib
import argparse

ROOT = Path(__file__).resolve().parents[2]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--template', type=Path, required=True, help='Original supplied Microsoft PowerPoint template')
args = parser.parse_args()
SOURCE = args.template.expanduser().resolve()
DEST = ROOT / 'docs/workshop/OpsPilot_MSA_Workshop.pptx'
SCREEN = ROOT / 'docs/opspilot-dashboard.png'
prs = Presentation(SOURCE)
original_ids = list(prs.slides._sldIdLst)
original_title = prs.slides[1]
avatar = next(s.image.blob for s in original_title.shapes if s.shape_type == 13)
layouts = prs.slide_masters[0].slide_layouts
BLUE='0078D4'; PURPLE='5C2D91'; INK='202D3D'; MUTED='526175'; PALE='EEF5FC'; LILAC='F3EFF9'; LINE='CFDDEB'; WHITE='FFFFFF'
notes=[]

def rgb(hex): return RGBColor.from_string(hex)
def txt(slide,x,y,w,h,text,size=22,bold=False,color=INK,font='Segoe UI',fill=None):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    if fill:shape.fill.solid();shape.fill.fore_color.rgb=rgb(fill)
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
        p.space_after=Pt(8);p.line_spacing=1.08
    return shape

def box(slide,x,y,w,h,fill=PALE,line=None):
    s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    s.adjustments[0]=.1;s.fill.solid();s.fill.fore_color.rgb=rgb(fill)
    if line:s.line.color.rgb=rgb(line)
    else:s.line.fill.background()
    return s

def label(slide,x,y,text,w=11.9,color=MUTED): return txt(slide,x,y,w,.3,text,12,color=color)
def title(slide,text):
    for sh in list(slide.shapes):
        if sh.is_placeholder: sh._element.getparent().remove(sh._element)
    txt(slide,.65,.47,12.02,.75,text,32,True)

def new(name,minutes,body_notes,layout=6):
    s=prs.slides.add_slide(layouts[layout]);title(s,name)
    index=len(notes)+1
    label(s,.65,7.04,'OpsPilot workshop | Adarsh Gupta',w=10)
    label(s,12.1,7.04,f'{index:02d}',w=.6)
    n=f'{name}\nSuggested time: {minutes} min\n\n{body_notes}'
    s.notes_slide.notes_text_frame.text=n
    notes.append((name,minutes,body_notes))
    return s

def points(slide,x,y,w,items,size=23,step=.95):
    for i,(head,body) in enumerate(items):
        txt(slide,x,y+i*step,w,.38,head,size,True)
        txt(slide,x,y+i*step+.43,w,step-.45,body,18,color=MUTED)

def card(slide,x,y,w,h,head,body,accent=BLUE):
    box(slide,x,y,w,h)
    txt(slide,x+.22,y+.2,w-.44,.5,head,23,True,accent)
    txt(slide,x+.22,y+.88,w-.44,h-1.03,body,20)

def code(slide,x,y,w,h,text,size=17):
    box(slide,x,y,w,h,fill='182D48')
    shape=txt(slide,x+.22,y+.2,w-.44,h-.35,text,size,color='EDF5FF',font='Consolas')
    for p in shape.text_frame.paragraphs:
        p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(size*1.25)
    return shape

def source(slide,labeltext,url):
    shape=txt(slide,.65,6.32,12,.23,labeltext,10,color=MUTED)
    shape.text_frame.paragraphs[0].runs[0].hyperlink.address=url

def arrow(slide,x1,y1,x2,y2):
    line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    line.line.color.rgb=rgb(BLUE);line.line.width=Pt(2)
    tail=OxmlElement('a:tailEnd');tail.set('type','triangle');line.line._get_or_add_ln().append(tail)

def shot(slide,x,y,w,h,region):
    # Native PowerPoint cropping: preserve the original screenshot and editable picture.
    shape=slide.shapes.add_picture(str(SCREEN),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
    left,top,right,bottom=region
    shape.crop_left=left;shape.crop_top=top;shape.crop_right=right;shape.crop_bottom=bottom
    return shape

# 1: supplied walk-in layout and presenter portrait.
s=new('OpsPilot',1,'Welcome students. Introduce yourself as Adarsh Gupta, Beta Student Ambassador, retaining the designation from the supplied deck. Explain that today we will explore a working AI incident copilot and modify fictional inputs. All remediation is simulated. Start the 60-minute session clock.',4)
for sh in list(s.shapes):
    if sh.has_text_frame and sh.text=='OpsPilot':sh._element.getparent().remove(sh._element)
txt(s,.65,2.0,7.9,.95,'OpsPilot',52,True)
txt(s,.67,3.0,7.7,.65,'Build an AI Incident Copilot',30)
txt(s,.67,3.82,7.6,.7,'Adarsh Gupta\nBeta Student Ambassador',19,color=MUTED)
txt(s,.67,4.77,7,.4,'.NET 8 + Microsoft Foundry | 60-minute workshop',16,color=PURPLE)
s.shapes.add_picture(BytesIO(avatar),Inches(9.15),Inches(1.25),width=Inches(1.65),height=Inches(1.65))

s=new('What you will be able to do',2,'Audience: students with basic C# familiarity. Azure experience is helpful but not required. Students without model access can follow the instructor demo and run the offline tests. This is a code walkthrough and modification workshop, not a claim that everyone will write the full application from scratch in one hour. Run of show: 5 minutes context, 10 minutes tour, 15 minutes implementation, 15 minutes demo and practice, 10 minutes review, 5 minutes wrap-up.')
points(s,.7,1.58,6.5,[('Connect a .NET app to a model','Use an Azure identity and IChatClient.'),('Turn logs into a typed assessment','Ask for a schema, then validate the response.'),('Keep people in control','Review evidence before approving a simulation.'),('Test both success and failure','Explore fictional incidents and safe fallbacks.')],step=1.05)
box(s,8.0,1.55,4.6,4.55,fill=LILAC)
txt(s,8.3,1.82,4,.4,'Our 60 minutes',24,True,PURPLE)
txt(s,8.3,2.45,4,3.3,'05 min   Frame the problem\n10 min   Tour OpsPilot\n15 min   Read the core code\n15 min   Demo + hands-on\n10 min   Review and test\n05 min   Wrap up',21)

s=new('Checkout is failing. What do you know?',2,'Ask for observations before explanations. The logs show a new deployment, HTTP 500s, a TaxMapper exception and 82 failures in 100 checkout requests in a 60-second window. The release is correlated with the error onset; it is not proof that the release is the root cause. Ask: what would you inspect next? All data is fictional. The displayed log text is shortened for teaching; the application receives the full fixture.')
code(s,.7,1.6,11.95,2.75,'10:00  INFO   deployment version=v2.4 previous=v2.3\n10:01  ERROR  POST /checkout HTTP 500\n10:01  ERROR  NullReferenceException at TaxMapper.Map(...)\n10:02  WARN   checkout_requests=100 http_500=82 window=60s',22)
card(s,.7,4.7,5.8,1.35,'Observation','What does the log actually say?')
card(s,6.85,4.7,5.8,1.35,'Hypothesis','What might explain the failure?',PURPLE)

s=new('Meet OpsPilot',3,'Show the actual dashboard. It is a Blazor Web App using Interactive Server rendering. Select the checkout scenario, point out editable logs, and explain that Azure receives only the submitted fictional incident data and fixture context. The screenshot is from this project, not a mock product image. Live actions will be shown again during the demo. Mention that output wording varies across calls.')
points(s,.7,1.65,3.5,[('Choose an incident','Use a sample or fictional logs.'),('Review an assessment','Inspect evidence and causes.'),('Share the update','Copy a stakeholder update.'),('Approve a simulation','Create a local text record only.')],size=21,step=1.08)
shot(s,4.65,1.55,7.95,4.65,(0,0,0,.55))

s=new('One request, from browser to assessment',3,'Walk through the diagram left to right. Blazor component events execute on the server and call the analyzer through dependency injection; there is no custom REST analysis endpoint. The analyzer enforces input and analysis limits, reads JSON fixtures, then sends a strict system message and an untrusted data message through IChatClient. AzureOpenAIClient authenticates through DefaultAzureCredential. Typed JSON is validated before display. Failure returns a conservative response. No function calling, RAG, or multi-agent loop is enabled.')
for x,h,b in [(.7,'Blazor UI','Fictional input'),(3.2,'Analyzer','Validate + add context'),(5.7,'IChatClient','AzureOpenAIClient'),(8.2,'Foundry','Structured JSON')]:
    box(s,x,1.8,2.2,1.45);txt(s,x+.13,2.02,1.95,.4,h,21,True,BLUE);txt(s,x+.13,2.6,1.96,.5,b,16)
for x in [2.9,5.4,7.9]:arrow(s,x,2.52,x+.3,2.52)
box(s,10.85,1.8,1.8,1.45,fill=LILAC);txt(s,11.0,2.0,1.5,.45,'Validate',20,True,PURPLE);txt(s,11.0,2.6,1.5,.5,'Show result',16)
arrow(s,10.4,2.52,10.85,2.52)
card(s,.7,4.0,3.7,1.8,'Local fixtures','Runbooks + deployment history')
card(s,4.82,4.0,3.7,1.8,'Azure identity','CLI login locally; managed identity in Azure')
card(s,8.95,4.0,3.7,1.8,'Failure path','Safe fallback + human review',PURPLE)

s=new('The small stack behind the demo',2,'Keep this slide about each technology\'s job. Bootstrap is a third-party CSS library, so do not reuse the old deck\'s claim of no third-party packages. The app targets .NET 8. Microsoft.Extensions.AI and its OpenAI adapter provide the abstraction; Azure.AI.OpenAI supplies AzureOpenAIClient; Azure.Identity provides DefaultAzureCredential. Dependency versions are pinned in the project and lock files. The only custom JavaScript is clipboard interop.')
for x,y,h,b in [(.7,1.55,'.NET 8 + ASP.NET Core','Hosting, dependency injection and health checks'),(6.8,1.55,'Blazor + Bootstrap','Interactive Server UI; no JavaScript UI framework'),(.7,3.76,'Microsoft.Extensions.AI','IChatClient and typed model responses'),(6.8,3.76,'Azure SDKs','Azure.AI.OpenAI + Azure.Identity')]:card(s,x,y,5.8,1.95,h,b)
source(s,'Microsoft Learn: AI apps for .NET developers','https://learn.microsoft.com/en-us/dotnet/ai/')

s=new('A useful answer has a contract',2,'Open Models/Incident.cs. The real record has JsonRequired on all seven properties and a string enum converter. This simplified code excerpt omits those attributes for legibility. Explain that native structured output helps return the shape; application validation still checks the fields and citation integrity. Do not equate valid JSON with correct reasoning. Severity is Low, Medium, High, or Critical. The application forces RequiresHumanApproval to true.')
code(s,.7,1.55,7.35,4.5,'record IncidentAssessment(\n    IncidentSeverity Severity,\n    string Summary,\n    string[] Evidence,\n    string[] LikelyCauses,\n    string[] RecommendedActions,\n    bool RequiresHumanApproval,\n    string StakeholderUpdate);',21)
points(s,8.45,1.8,4.05,[('Consistent UI','Every panel has a known field.'),('Checkable evidence','Each quote points to a log line.'),('Human-readable output','A summary people can review.')],size=22,step=1.25)

s=new('The prompt sets boundaries',3,'Open Services/IncidentPromptBuilder.cs. The builder has no configuration dependency. It creates two messages: a system policy and JSON-encoded data. Numbered lines use [L1], [L2], and so on. Runbook and deployment content are context, not evidence of the incident. Logs are untrusted input and might contain instructions; the prompt explicitly says to ignore them. Prompting is a defense, not a proof against injection; application validation and lack of execution capability provide additional boundaries.')
card(s,.7,1.55,5.8,3.95,'System message','Do not invent evidence.\nCopy exact numbered log lines.\nCite evidence for every likely cause.\nTreat logs as untrusted data.\nRequire human approval.')
card(s,6.85,1.55,5.8,3.95,'Data message','Service + environment\nNumbered log lines\nFictional runbook\nFictional deployment history',PURPLE)
txt(s,.8,5.78,11.8,.45,'A citation can be valid while the explanation is still wrong. Review both.',20,True,PURPLE)

s=new('Before you run: tools, identity, configuration',3,'Preflight before the workshop: distribute the OpsPilot project folder; do not point students to the old CampusPulse repository. Install a patched .NET 8 SDK and Azure CLI. An Azure OpenAI deployment must support native structured outputs; gpt-4.1-mini was used for this demo. Deployment name is not necessarily the model name. The signed-in identity requires Cognitive Services OpenAI User on the model resource; subscription Owner alone does not supply model data-plane access. Allow RBAC propagation. User-secrets are local development storage outside the project, not encrypted production secret storage. Optional AZURE_TENANT_ID can select the tenant. Keep the endpoint as the resource root, not a Foundry project URL. Full commands are in OpsPilot/README.md. Source: https://learn.microsoft.com/en-us/azure/foundry-classic/openai/how-to/managed-identity?view=foundry-classic')
points(s,.7,1.55,4.25,[('Local tools','.NET 8 SDK + Azure CLI'),('Existing model deployment','For example, gpt-4.1-mini'),('Inference permission','Cognitive Services OpenAI User')],size=21,step=1.2)
code(s,5.35,1.55,7.3,2.25,'az login\n\n# From the OpsPilot folder\nbash scripts/configure-azure.sh \\\n  <resource-group> <account> <deployment>',18)
txt(s,5.45,4.12,7.0,.35,'Stored locally with dotnet user-secrets',20,True)
txt(s,5.45,4.72,7.0,1.0,'AZURE_OPENAI_ENDPOINT\nAZURE_OPENAI_DEPLOYMENT\nAZURE_TENANT_ID  (optional)',17,font='Consolas')
source(s,'Setup commands and role guidance: OpsPilot/README.md','OpsPilot/README.md')

s=new('Connect the model through IChatClient',3,'Open Services/AzureChatClientFactory.cs. This teaching excerpt leaves out endpoint checks, optional TenantId, disabled SDK content logging, network timeout and retry settings; those remain in the actual implementation. DefaultAzureCredential uses a credential chain, which can use the current Azure CLI identity locally and a managed identity when hosted. No API key is embedded in code. The factory returns the Azure chat client as IChatClient and Program.cs registers it with DI. Source: https://learn.microsoft.com/en-us/dotnet/ai/quickstarts/structured-output')
code(s,.7,1.55,11.95,2.95,'var azureClient = new AzureOpenAIClient(\n    new Uri(endpoint),\n    new DefaultAzureCredential());\n\nIChatClient chat = azureClient\n    .GetChatClient(deployment).AsIChatClient();',22)
card(s,.7,4.65,5.8,1.4,'Locally','Use your Azure CLI identity.')
card(s,6.85,4.65,5.8,1.4,'When hosted','Use the app\'s managed identity.',PURPLE)
source(s,'Microsoft Learn: request structured output with .NET','https://learn.microsoft.com/en-us/dotnet/ai/quickstarts/structured-output')

s=new('Request a type. Then validate it.',3,'Open Services/FoundryIncidentAnalyzer.cs and AssessmentSafety.cs. This is a shortened excerpt; the real analyzer also supplies serializer options and a token budget. Explain the cancellation token. Native JSON schema comes from IncidentAssessment. The validator requires real enum values, bounded nonempty fields, exact evidence quotes, and cause references pointing to included evidence. It rejects invalid responses with a safe fallback. Show the final override that makes RequiresHumanApproval true. This is one model request, not an autonomous agent workflow.')
code(s,.7,1.55,11.95,2.75,'var response = await chat.GetResponseAsync<IncidentAssessment>(\n    messages,\n    useJsonSchemaResponseFormat: true,\n    cancellationToken: cancellationToken);\n\nvar assessment = response.Result;',22)
for x,h,b in [(.7,'1  Shape','Required, bounded fields'),(4.8,'2  Evidence','Exact log quotes + citations'),(8.9,'3  Approval','Always enforced by the app')]:card(s,x,4.65,3.75,1.6,h,b)
source(s,'Microsoft Learn: structured outputs','https://learn.microsoft.com/en-us/azure/foundry/openai/how-to/structured-outputs')

s=new('Give the model useful context',3,'Open Fixtures/runbooks.json and deployments.json, then Services/Contracts.cs. The analyzer calls both services directly before inference. The interfaces are read-only and accept cancellation tokens. Unknown services have no fixture context; this is an intentional behavior, not a database lookup failure. The same methods could be wrapped as Foundry function tools in a future workshop. This project does not register tools, use retrieval, or query real deployment history.')
code(s,.7,1.6,11.95,1.85,'GetRunbookAsync(serviceName, cancellationToken)\n\nGetRecentDeploymentsAsync(serviceName, cancellationToken)',23)
card(s,.7,3.85,5.8,2.0,'Runbook','Investigation steps and review guidance for each service.')
card(s,6.85,3.85,5.8,2.0,'Deployment history','Fictional releases, timestamps and descriptions.',PURPLE)
label(s,.8,6.03,'Today: direct service calls. Later: possible function-tool wrappers.')

s=new('Let\'s investigate a live demo',.5,'Switch to the browser at http://localhost:5188. Keep the terminal visible nearby. Start the app in advance with dotnet run --project src/OpsPilot.Web from the OpsPilot folder. If Azure is unavailable, show the saved dashboard screenshot and demonstrate the real fallback without pretending a live assessment succeeded.',19)
# Respect the original section-divider layout instead of forcing a normal title.
for sh in list(s.shapes):
    if sh.has_text_frame and sh.text=="Let's investigate a live demo":sh._element.getparent().remove(sh._element)
txt(s,.65,2.9,10.2,.75,'Live demo: investigate checkout',36)
txt(s,.67,3.9,9.5,.45,'http://localhost:5188',23,color=BLUE)
txt(s,.67,4.63,8.9,.65,'Observe. Explain. Review. Simulate.',24,color=PURPLE)

s=new('The demo: from log to local request',4.5,'Demo steps: 1. Select Checkout fails after release. 2. Click Analyze incident and show loading/cancel. 3. Trace a likely cause back to exact evidence. 4. Read the stakeholder update and copy it. 5. Point to the predetermined demo record, then click Approve remediation. 6. Show the session activity and disabled duplicate approval. 7. Change the environment field and show that the old assessment disappears. Explain that the approval is bound to the assessed input. The model does not execute a rollback; the record is generated by local C# code. Source: Components/Pages/Home.razor.cs and Services/RemediationSimulation.cs.')
points(s,.7,1.6,5.15,[('Select checkout-api','Analyze the sample deployment failure.'),('Trace a cause to its evidence','Confirm that the cited line supports the claim.'),('Copy the stakeholder update','Separate facts from hypotheses.'),('Approve remediation','Only a local simulation record is created.')],size=21,step=1.05)
box(s,6.35,1.55,6.3,4.45,fill=LILAC)
txt(s,6.65,1.87,5.7,.5,'The only remediation outcome',23,True,PURPLE)
code(s,6.65,2.7,5.7,1.55,'Rollback request created\nfor checkout-api-v2.4.',21)
txt(s,6.65,4.62,5.65,.85,'No rollback. No infrastructure access.\nNo external request is sent.',22,True)

s=new('Your turn: change the evidence',10,'Run this ten-minute paired exercise. Minute 0-2: open the shared OpsPilot folder, start the app, select one scenario and predict the likely issue before calling the model. Minute 2-5: analyze it and verify every cause has a citation supporting it. Minute 5-8: remove the most diagnostic log line (TaxMapper, pool saturation, or issuer mismatch) and analyze again; compare uncertainty rather than expecting exact text. Minute 8-10: copy the update and explain to a partner which recommendation needs review. Students with no Azure access should pair with a configured machine, inspect the fixtures and prompt, and run the offline tests. Do not share personal access tokens or keys. Success: one supported observation, one hypothesis and a justified next check. Do not add real data.')
code(s,.7,1.5,11.95,1.28,'cd OpsPilot\ndotnet run --project src/OpsPilot.Web',22)
for x,h,b in [(.7,'Predict','Choose a sample.\nName the first clue.'),(4.8,'Change','Remove one key log line.\nAnalyze again.'),(8.9,'Explain','Compare the evidence.\nWhat is less certain?')]:card(s,x,3.2,3.75,2.2,h,b)
box(s,.7,5.72,11.95,.45,fill=LILAC);txt(s,.85,5.77,11.6,.33,'Done when you can separate an observation, a hypothesis and a next check.',18,True,PURPLE)

s=new('Three incidents, three different clues',3,'Invite a pair to report from each scenario. Database: saturated pool, long stock-reservation transaction and timeouts; reachable database means connectivity alone does not explain the wait. Checkout: release timing and a TaxMapper exception; timing alone is insufficient. Authentication: expected and actual issuer differ even though the signature is valid and token not expired. Never solve the authentication case by disabling validation. These are instructor expectations from fictional fixtures, not fixed model outputs.')
for x,h,b in [(.7,'PostgreSQL pool','Clue: active=100, idle=0\nSymptom: waits + HTTP 503\nNext check: transaction lifetime'),(4.8,'Checkout release','Clue: TaxMapper exception\nSymptom: HTTP 500\nNext check: release change'),(8.9,'Authentication','Clue: issuer mismatch\nSymptom: HTTP 401\nNext check: intended issuer')]:card(s,x,1.65,3.75,3.8,h,b)
txt(s,.8,5.8,11.9,.45,'Which missing log line changed your confidence the most?',24,True,PURPLE)

s=new('A safe demo must handle bad answers',3,'Explain the safeguards by behavior. Input validation bounds identifiers, logs and line count. The model receives strict instructions, but the app also validates the output. Evidence checks establish quote and citation integrity; they do not prove reasoning is correct. A failed request, invalid JSON, invalid evidence or context failure returns fallback. Its High enum is explicitly labeled a conservative review priority, not a measured incident severity. The UI disables approval on fallback. Every assessment requires human approval; model text never becomes a command. App authentication is intentionally out of scope, so this is a controlled workshop demo.')
points(s,.7,1.6,5.6,[('Before the request','Validate input and apply shared limits.'),('After the response','Check schema, exact quotes and citations.'),('Before approval','Require a grounded result and explicit review.')],size=22,step=1.28)
box(s,7.0,1.6,5.65,4.3,fill=LILAC)
txt(s,7.3,1.92,5.05,.5,'When analysis fails',26,True,PURPLE)
txt(s,7.3,2.72,5.05,2.2,'Say assessment is unavailable.\nInvent no evidence or causes.\nRequire human review.\nDisable simulated approval.',22)
txt(s,7.3,5.13,5.05,.55,'A fallback is better than false certainty.',19,True,PURPLE)

s=new('How do we know it works?',2,'Run dotnet test from the OpsPilot folder. Offline tests use deterministic fake IChatClient responses and do not need Azure. They cover prompt boundaries, malformed/missing fields, evidence integrity, caller cancellation, timeout, fixture lookup, request limits, simulation and safe exception handling. The live Azure test is deliberately opt-in, exercises all three scenarios and incurs model usage. It must fail on fallback, not silently pass. Browser verification covers interaction, copy, approval, input invalidation and mobile layout. GET /health is liveness only, never proof that Azure access works. Use current commands rather than promising a fixed test count forever.')
code(s,.7,1.6,11.95,2.35,'dotnet build OpsPilot.sln\ndotnet test OpsPilot.sln\n\n# Opt-in live inference (Bash; uses Azure model tokens)\nOPSPILOT_LIVE_TESTS=1 dotnet test --filter Category=LiveAzure',19)
card(s,.7,4.4,3.75,1.7,'Unit tests','Prompt + fallback + validation')
card(s,4.8,4.4,3.75,1.7,'Live tests','Actual model + all scenarios')
card(s,8.9,4.4,3.75,1.7,'Browser checks','Copy + approval + cancellation')

s=new('If the workshop hits a snag',2,'Use metadata-only logs, not raw exceptions, credentials or user log content. 401/403: verify Azure CLI account and tenant, and the inference role on the selected model resource; newly assigned roles can take time to propagate. In this implementation session propagation took about ten minutes, but do not promise a fixed delay. 400: verify endpoint, deployment and native structured-output support. 429: distinguish Azure quota from the app\'s process-wide 10 attempts/minute and 2 concurrent analyses. Pair students or stagger requests. Copy needs localhost or HTTPS. If the model produces invalid evidence, show fallback and review the input before retrying. Source: OpsPilot/README.md. Publish guidance uses managed identity, HTTPS and session affinity; no deployment is required during this workshop.')
rows=[('401 / 403','Check identity, tenant and inference role.'),('400 / no assessment','Check endpoint, deployment and schema support.'),('429 / demo busy','Wait, then stagger or pair requests.'),('Clipboard unavailable','Use localhost / HTTPS, or copy the text manually.'),('Healthy, but AI fails','Health checks liveness; analyze a sample to test AI.')]
for i,(a,b) in enumerate(rows):
    y=1.65+i*.86
    box(s,.7,y,11.95,.7,fill=PALE if i%2==0 else 'F7F8FA')
    txt(s,.9,y+.15,3.45,.4,a,20,True,BLUE);txt(s,4.6,y+.15,7.8,.4,b,19)

s=new('What you built, and where it stops',2,'Ask students to summarize the pattern in their own words. We have a UI, typed request/response contract, contextual model call, validation and a human approval boundary. We do not have a production incident-management platform. There is no real rollback, Kubernetes, GitHub, Azure Monitor, authentication, database persistence, RAG, Azure AI Search or multi-agent workflow. Possible follow-up exercise: add another fictional JSON scenario and a prompt test; keep scope small. Hosting guidance is in the README and uses managed identity. Use https://github.com/Adarsh1999/OpsPilot and arrange participant access according to repository visibility.')
card(s,.7,1.65,5.8,3.85,'The reusable pattern','Clear contract\nUseful context\nStructured response\nApplication validation\nHuman review')
card(s,6.85,1.65,5.8,3.85,'The deliberate boundary','Fictional logs and fixtures\nLocal session activity\nNo real infrastructure actions\nNo autonomous agent loop',PURPLE)
txt(s,.8,5.84,11.8,.4,'Next challenge: add a fourth fictional scenario and one meaningful test.',21,True)

s=new('Keep learning with Microsoft Learn',2,'All links on this slide are clickable. Share the existing OpsPilot folder plus its README; use https://github.com/Adarsh1999/OpsPilot and check participant access before the workshop. The .NET AI hub and structured-output quickstart support the code pattern in this workshop. The Azure identity article explains keyless model access. The Student Hub offers learning and community opportunities; avoid promising admission, benefits or program requirements. Sources were checked while preparing the deck.')
links=[('Build AI apps with .NET','https://learn.microsoft.com/en-us/dotnet/ai/'),('Request structured output in C#','https://learn.microsoft.com/en-us/dotnet/ai/quickstarts/structured-output'),('Connect to Azure OpenAI with Microsoft Entra ID','https://learn.microsoft.com/en-us/azure/foundry-classic/openai/how-to/managed-identity?view=foundry-classic'),('Explore the Microsoft Learn Student Hub','https://learn.microsoft.com/en-us/training/student-hub/become-a-student-ambassador')]
for i,(head,url) in enumerate(links):
    y=1.6+i*1.05
    shape=txt(s,.8,y,11.8,.43,head,24,True,BLUE);shape.text_frame.paragraphs[0].runs[0].hyperlink.address=url
    txt(s,.8,y+.51,11.8,.3,{'Build AI apps with .NET':'learn.microsoft.com/dotnet/ai','Request structured output in C#':'Microsoft Learn / .NET AI / Structured output','Connect to Azure OpenAI with Microsoft Entra ID':'Microsoft Learn / Azure OpenAI / Managed identity','Explore the Microsoft Learn Student Hub':'Microsoft Learn / Student Hub'}[head],15,color=MUTED)

s=new('Thank you',1,'Close with a quick retrieval question: what is the difference between a model returning valid JSON and a trustworthy incident assessment? Invite students to name one evidence check and one approval boundary. Ask for questions. Thank the participants and direct them to Microsoft Learn and the shared project README. Keep this final slide visible while students finish.',28)
for sh in list(s.shapes):
    if sh.has_text_frame and sh.text=='Thank you':sh._element.getparent().remove(sh._element)
txt(s,.67,2.25,8.8,.85,'Thank you!',44,True)
txt(s,.67,3.35,8.8,.65,'What would you verify before acting?',28,color=PURPLE)
txt(s,.67,4.4,8.5,.5,'Questions + discussion',24)
txt(s,.67,5.18,7.5,.5,'Adarsh Gupta | Beta Student Ambassador',18,color=MUTED)

# Drop original CampusPulse slides only after adding replacements, keeping package names unique.
for sid in original_ids:
    prs.part.drop_rel(sid.rId)
    prs.slides._sldIdLst.remove(sid)
prs.core_properties.title='OpsPilot - AI Incident Copilot Workshop'
prs.core_properties.subject='Microsoft Learn Student Ambassadors | .NET 8 + Microsoft Foundry'
prs.core_properties.author='Adarsh Gupta'
prs.core_properties.keywords='OpsPilot, Microsoft Learn, Student Ambassadors, .NET 8, Azure, workshop'
prs.core_properties.comments='Adapted from the supplied Microsoft Student Ambassadors template. Fictional incident workshop.'
buffer=BytesIO();prs.save(buffer)
# Preserve original theme, masters and layouts byte-for-byte, including their relationships.
with ZipFile(SOURCE) as original, ZipFile(buffer) as generated, ZipFile(DEST,'w',ZIP_DEFLATED) as out:
    preserve=[n for n in original.namelist() if n.startswith(('ppt/theme/','ppt/slideMasters/','ppt/slideLayouts/'))]
    for item in generated.infolist():
        out.writestr(item,original.read(item.filename) if item.filename in preserve else generated.read(item.filename))
    assert all(n in generated.namelist() for n in preserve)

md=['# OpsPilot workshop presenter guide','', 'Duration: 60 minutes. Audience: students with basic C# familiarity.','',
'## Before the session','', '- Share the OpsPilot folder and README with participants. Repository: https://github.com/Adarsh1999/OpsPilot; arrange participant access before the workshop.',
'- Test the configured Azure identity with an actual sample; `/health` checks liveness only.',
'- Start the app at http://localhost:5188. Keep the saved dashboard screenshot ready for network issues.',
'- Have participants work in pairs where Azure access is limited. Use only fictional data.',
'- Speaker notes are also embedded in each PowerPoint slide.','']
for i,(name,minutes,note) in enumerate(notes,1):md += [f'## {i:02d}. {name} ({minutes:g} min)','',note,'']
(ROOT/'docs/workshop/Presenter_Guide.md').write_text('\n'.join(md))
assert sum(n[1] for n in notes)==60
with ZipFile(SOURCE) as a, ZipFile(DEST) as b:
    assert all(a.read(n)==b.read(n) for n in preserve)
print(f'Created {DEST.name}: {len(notes)} slides, 60 minutes; {len(preserve)} original template parts preserved byte-for-byte.')
print('Original SHA256:',hashlib.sha256(SOURCE.read_bytes()).hexdigest())
